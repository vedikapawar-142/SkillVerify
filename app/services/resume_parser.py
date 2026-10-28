"""Resume text extraction, validation, and NLP-based skill extraction service.
Supports:
- PDF extraction via PyPDF2 / pypdf
- DOCX extraction via python-docx
- Binary magic bytes validation (%PDF- and PK\x03\x04)
- Configurable skills taxonomy dictionary
- Graceful handling of scanned, unreadable, or empty files
"""
import os
import re
import json
from typing import Dict, Any, List, Tuple

# Import PyPDF2 / pypdf
try:
    import PyPDF2 as pypdf
except ImportError:
    try:
        import pypdf
    except ImportError:
        pypdf = None

# Import python-docx
try:
    import docx
except ImportError:
    docx = None


SKILL_TAXONOMY = {
    "Programming Languages": [
        "Python", "Java", "C", "C++", "C#", "JavaScript", "TypeScript", 
        "Go", "Rust", "PHP", "Ruby", "Kotlin", "Swift", "Dart", "R", "Scala", "SQL"
    ],
    "Web Development": [
        "HTML", "HTML5", "CSS", "CSS3", "JavaScript", "React", "React.js", "Vue.js", 
        "Angular", "Node.js", "Express.js", "Flask", "Django", "FastAPI", 
        "Bootstrap", "Tailwind CSS", "REST API", "RESTful APIs", "GraphQL"
    ],
    "Databases": [
        "MySQL", "PostgreSQL", "SQLite", "MongoDB", "Redis", "Oracle", 
        "SQL Server", "Firebase", "Cassandra", "MariaDB"
    ],
    "AI and Machine Learning": [
        "Machine Learning", "Deep Learning", "Natural Language Processing", "NLP", 
        "Computer Vision", "PyTorch", "TensorFlow", "Keras", "Scikit-Learn", 
        "Pandas", "NumPy", "OpenCV", "Generative AI", "LLM", "Data Science",
        "Matplotlib", "Seaborn", "NLTK"
    ],
    "Cloud & DevOps": [
        "AWS", "Microsoft Azure", "Azure", "Google Cloud", "GCP", 
        "Docker", "Kubernetes", "CI/CD", "Terraform", "Serverless", "Linux"
    ],
    "Development Tools": [
        "Git", "GitHub", "GitLab", "VS Code", "Postman", "Linux", 
        "Ubuntu", "Jira", "PyCharm", "Bash", "Shell Scripting"
    ]
}


def validate_resume_file(file_path: str, max_size_bytes: int = 5 * 1024 * 1024) -> Tuple[bool, str]:
    """Validate that the file exists, is non-empty, within size limits,
    and has valid file extension and matching magic header bytes.
    """
    if not os.path.exists(file_path):
        return False, "File does not exist."

    file_size = os.path.getsize(file_path)
    if file_size == 0:
        return False, "The uploaded file is empty (0 bytes). Please upload a valid resume."

    if file_size > max_size_bytes:
        max_mb = max_size_bytes / (1024 * 1024)
        return False, f"File size ({file_size / (1024*1024):.1f} MB) exceeds the configured limit of {max_mb:.0f} MB."

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in (".pdf", ".docx"):
        return False, "Unsupported file format. Please upload a PDF or DOCX file."

    # Inspect magic bytes
    try:
        with open(file_path, "rb") as f:
            header = f.read(8)

        if ext == ".pdf":
            if not header.startswith(b"%PDF-"):
                return False, "Invalid PDF file: Missing standard PDF header signature."
        elif ext == ".docx":
            # DOCX is a zip archive starting with PK\x03\x04
            if not header.startswith(b"PK\x03\x04"):
                return False, "Invalid DOCX file: Missing valid Word document container signature."
    except Exception as e:
        return False, f"Error validating file signature: {str(e)}"

    return True, "Valid"


def extract_text_from_pdf(file_path: str) -> str:
    """Extract plain text from PDF using PyPDF2 / pypdf."""
    if not pypdf:
        raise RuntimeError("PyPDF2 / pypdf library is not installed.")
    
    text_content = []
    try:
        reader = pypdf.PdfReader(file_path)
        if len(reader.pages) == 0:
            raise ValueError("PDF document contains 0 pages.")

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    except Exception as e:
        raise ValueError(f"Could not read PDF document: {str(e)}")

    extracted = "\n".join(text_content).strip()
    if not extracted or len(extracted.strip()) < 20:
        raise ValueError(
            "Could not extract readable text from the PDF. The document may be a scanned image or empty. "
            "Please upload a standard text-based PDF or DOCX."
        )

    return extracted


def extract_text_from_docx(file_path: str) -> str:
    """Extract plain text from DOCX using python-docx."""
    if not docx:
        raise RuntimeError("python-docx library is not installed.")

    try:
        doc = docx.Document(file_path)
        text_content = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_content.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))
        extracted = "\n".join(text_content).strip()
    except Exception as e:
        raise ValueError(f"Could not read DOCX document: {str(e)}")

    if not extracted or len(extracted.strip()) < 20:
        raise ValueError("Could not extract readable text from the DOCX file. The document is empty.")

    return extracted


def extract_resume_text(file_path: str) -> str:
    """Route file to appropriate text extractor based on extension."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Resume file not found at: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF or DOCX file.")


def extract_contact_info(text: str) -> Dict[str, str]:
    """Extract candidate name, email, and phone."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    # Email regex
    email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
    email = email_match.group(0).lower() if email_match else ""

    # Phone regex
    phone_match = re.search(r"(?:(?:\+|0{0,2})91[\s-]?)?[6789]\d{9}|(?:\+?1[\s-]?)?\(?\d{3}\)?[\s-]?\d{3}[\s-]?\d{4}", text)
    phone = phone_match.group(0) if phone_match else ""

    # Name heuristics
    candidate_name = ""
    for line in lines[:6]:
        cleaned = re.sub(r"[^a-zA-Z\s]", "", line).strip()
        lower_line = line.lower()
        if any(w in lower_line for w in ["resume", "curriculum", "vitae", "email", "phone", "contact", "@", "page"]):
            continue
        if 2 <= len(cleaned.split()) <= 4 and len(cleaned) < 40:
            candidate_name = cleaned.title()
            break

    if not candidate_name and lines:
        candidate_name = lines[0][:40].strip()

    return {
        "name": candidate_name,
        "email": email,
        "phone": phone
    }


def extract_sections(text: str) -> Dict[str, str]:
    """Split resume text into sections using regex headers."""
    section_patterns = {
        "education": r"(?:education|academic\s+background|qualifications|academic\s+details)",
        "experience": r"(?:work\s+experience|professional\s+experience|experience|employment\s+history)",
        "internships": r"(?:internships?|internship\s+experience|industrial\s+training)",
        "projects": r"(?:academic\s+projects|personal\s+projects|key\s+projects|projects)",
        "certifications": r"(?:certifications?|certificates?|licenses?|courses)",
        "skills": r"(?:technical\s+skills|skills|technologies|core\s+competencies|programming\s+languages)"
    }

    found_headers = []
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        clean_line = line.strip().lower()
        if len(clean_line) < 45:
            for sec_name, pat in section_patterns.items():
                if re.fullmatch(rf"[:\s\-_*#]*{pat}[:\s\-_*#]*", clean_line):
                    found_headers.append((idx, sec_name))
                    break

    sections = {k: "" for k in section_patterns}
    if not found_headers:
        sections["skills"] = text
        return sections

    found_headers.sort(key=lambda x: x[0])
    for i, (line_idx, sec_name) in enumerate(found_headers):
        end_idx = found_headers[i + 1][0] if i + 1 < len(found_headers) else len(lines)
        section_lines = lines[line_idx + 1:end_idx]
        sections[sec_name] = "\n".join(section_lines).strip()

    return sections


def analyze_resume_text(text: str) -> Dict[str, Any]:
    """Perform rule-based NLP extraction on resume text using configurable skills taxonomy."""
    contact = extract_contact_info(text)
    sections = extract_sections(text)

    categorized_skills: Dict[str, List[str]] = {}
    skill_evidence: Dict[str, str] = {}
    all_claimed_skills: List[str] = []

    text_lower = text.lower()

    for category, skill_list in SKILL_TAXONOMY.items():
        found_in_category = []
        for skill in skill_list:
            escaped = re.escape(skill.lower())
            if skill.lower() in ("c", "r"):
                # Strict boundary for single letters
                pattern = rf"(?:^|[\s,;|/])({escaped})(?:$|[\s,;|/])"
            elif skill.lower() in ("c++", "c#"):
                pattern = rf"(?:^|[\s,;|/])({escaped})(?:$|[\s,;|/])"
            else:
                pattern = rf"\b{escaped}\b"

            match = re.search(pattern, text_lower)
            if match:
                if skill not in found_in_category:
                    found_in_category.append(skill)
                if skill not in all_claimed_skills:
                    all_claimed_skills.append(skill)
                    start = max(0, match.start() - 20)
                    end = min(len(text), match.end() + 20)
                    skill_evidence[skill] = text[start:end].replace("\n", " ").strip()

        if found_in_category:
            categorized_skills[category] = sorted(found_in_category)

    # Extract structured section highlights
    education_text = sections.get("education", "")
    edu_snippets = []
    if education_text:
        edu_snippets = [l.strip() for l in education_text.splitlines() if len(l.strip()) > 3][:4]
    elif re.search(r"(?:diploma|b\.?tech|b\.?e\.?|bca|mca|b\.?sc|engineering)", text_lower):
        for line in text.splitlines():
            if re.search(r"(?:diploma|b\.?tech|b\.?e\.?|bca|mca|b\.?sc|engineering|university|college|institute)", line.lower()):
                edu_snippets.append(line.strip())
                if len(edu_snippets) >= 3:
                    break

    projects_text = sections.get("projects", "")
    project_snippets = []
    if projects_text:
        project_snippets = [l.strip() for l in projects_text.splitlines() if len(l.strip()) > 8][:5]

    exp_text = sections.get("experience", "") + "\n" + sections.get("internships", "")
    exp_snippets = []
    if exp_text.strip():
        exp_snippets = [l.strip() for l in exp_text.splitlines() if len(l.strip()) > 8][:5]

    cert_text = sections.get("certifications", "")
    cert_snippets = []
    if cert_text:
        cert_snippets = [l.strip() for l in cert_text.splitlines() if len(l.strip()) > 5][:4]

    return {
        "contact": contact,
        "claimed_skills": sorted(all_claimed_skills),
        "categorized_skills": categorized_skills,
        "skill_evidence": skill_evidence,
        "sections": {
            "education": edu_snippets,
            "projects": project_snippets,
            "experience": exp_snippets,
            "certifications": cert_snippets
        },
        "raw_text_length": len(text)
    }
