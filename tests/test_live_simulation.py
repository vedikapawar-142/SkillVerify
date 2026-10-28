"""Live end-to-end simulation of the core Phase 1 platform workflow.
Simulates candidate registration, resume upload & skill extraction,
7-question MCQ assessment taking & answer evaluation, answer review,
report & certificate generation, public verification, and HR dashboard inspection.
"""
import io
import docx
import pytest
from app import create_app
from app.extensions import db
from app.models import User, Attempt, Question
from config import TestingConfig
from app.services.assessment_service import get_or_seed_questions


def test_full_pipeline_simulation():
    app = create_app(TestingConfig)
    client = app.test_client()

    with app.app_context():
        # Ensure tables and questions are seeded
        get_or_seed_questions()

        # Step 1: Candidate registers and logs in
        reg_res = client.post("/register", data={
            "name": "Rahul Sharma",
            "email": "rahul@skillverify.com",
            "password": "cand123456"
        }, follow_redirects=True)
        assert reg_res.status_code == 200
        assert b"Registration successful" in reg_res.data or b"Dashboard" in reg_res.data

        login_res = client.post("/login", data={
            "email": "rahul@skillverify.com",
            "password": "cand123456"
        }, follow_redirects=True)
        assert login_res.status_code == 200
        assert b"Candidate Dashboard" in login_res.data or b"Rahul Sharma" in login_res.data

        # Step 2: Candidate creates and uploads realistic resume
        doc = docx.Document()
        doc.add_heading("Rahul Sharma", 0)
        doc.add_paragraph("Email: rahul@skillverify.com | Phone: 9876543210")
        doc.add_heading("Technical Skills", level=1)
        doc.add_paragraph("Python, Flask, MySQL, Git, Docker, REST API, Pandas, Data Structures")
        doc.add_heading("Education", level=1)
        doc.add_paragraph("B.Tech in Computer Science and Engineering, 2024")
        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)

        upload_res = client.post(
            "/resume/upload",
            data={"resume_file": (doc_bytes, "rahul_resume.docx")},
            content_type="multipart/form-data",
            follow_redirects=True
        )
        assert upload_res.status_code == 200
        assert b"Resume Analysis" in upload_res.data or b"Extracted Skills" in upload_res.data

        # Step 3: Candidate confirms claimed skills
        review_res = client.post(
            "/resume/review",
            data={"skills_list": "Python, Flask, MySQL, Git, Docker, REST API"},
            follow_redirects=True
        )
        assert review_res.status_code == 200

        # Step 4: Candidate initiates Python 7-question MCQ assessment
        take_page = client.get("/assessment/start/python", follow_redirects=True)
        assert take_page.status_code == 200
        assert b"Python Assessment" in take_page.data

        # Find the active attempt
        user = User.query.filter_by(email="rahul@skillverify.com").first()
        attempt = Attempt.query.filter_by(user_id=user.id, skill="python", status="in_progress").first()
        assert attempt is not None
        q_ids = attempt.get_question_ids()
        assert len(q_ids) == 7

        # Candidate answers all 7 questions correctly
        questions = Question.query.filter(Question.id.in_(q_ids)).all()
        submission = {f"q_{q.id}": q.correct_answer for q in questions}

        submit_res = client.post(f"/assessment/take/{attempt.id}", data=submission, follow_redirects=True)
        assert submit_res.status_code == 200
        assert b"Skill Assessment Result" in submit_res.data
        assert b"100" in submit_res.data
        assert b"Correct Answer" in submit_res.data
        assert b"Explanation" in submit_res.data

        # Step 5: Verify PDF Report and Certificate generation
        report = attempt.report
        assert report is not None
        report_res = client.get(f"/report/{report.id}/pdf")
        assert report_res.status_code == 200
        assert report_res.headers.get("Content-Type") == "application/pdf"

        cert = attempt.certificate
        assert cert is not None
        cert_res = client.get(f"/certificate/{cert.certificate_id}/download")
        assert cert_res.status_code == 200
        assert cert_res.headers.get("Content-Type") == "application/pdf"

        # Step 6: Public certificate verification
        verify_res = client.get(f"/certificate/verify/{cert.certificate_id}")
        assert verify_res.status_code == 200
        assert b"Authentic Certificate Verified" in verify_res.data

        # Step 7: HR recruiter logs in and views candidate record in HR dashboard
        client.get("/logout")
        hr_login = client.post("/hr/login", data={"email": "hr@skillverify.com", "password": "hr123456"}, follow_redirects=True)
        assert hr_login.status_code == 200
        assert b"Candidate Skill Verification Directory" in hr_login.data or b"HR Recruitment Management" in hr_login.data

        # Search for candidate
        search_res = client.get("/hr/dashboard?search=Rahul&lang=python&status=passed")
        assert search_res.status_code == 200
        assert b"Rahul Sharma" in search_res.data

        db.session.remove()
        db.drop_all()

