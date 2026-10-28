"""Comprehensive end-to-end and integration tests for Phase 1 of AI Skill Verification Platform."""
import io
import json
import pytest
import docx
from app import create_app
from app.extensions import db
from app.models import User, Resume, Question, Attempt, Report, Certificate
from config import TestingConfig
from app.services.resume_parser import analyze_resume_text
from app.services.assessment_service import (
    start_or_resume_attempt, evaluate_assessment, get_or_seed_questions,
    SUPPORTED_LANGUAGES
)


@pytest.fixture
def client():
    app = create_app(TestingConfig)
    with app.app_context():
        db.create_all()
        yield app.test_client()
        db.session.remove()
        db.drop_all()


def test_auth_candidate_registration_and_login(client):
    """Verify candidate registration, duplicate check, and login flow."""
    # Register candidate
    reg_res = client.post("/register", data={
        "name": "Arjun Patel",
        "email": "arjun@example.com",
        "password": "mypassword123"
    }, follow_redirects=True)
    assert reg_res.status_code == 200
    assert b"registration successful" in reg_res.data.lower()

    # Duplicate registration check
    dup_res = client.post("/register", data={
        "name": "Arjun Patel",
        "email": "arjun@example.com",
        "password": "mypassword123"
    }, follow_redirects=True)
    assert b"already exists" in dup_res.data

    # Login
    login_res = client.post("/login", data={
        "email": "arjun@example.com",
        "password": "mypassword123"
    }, follow_redirects=True)
    assert login_res.status_code == 200
    assert b"Candidate Dashboard" in login_res.data
    assert b"Arjun Patel" in login_res.data


def test_auth_hr_access_control(client):
    """Verify HR recruiter login and role-based access restrictions."""
    # Create HR via CLI/model
    app = client.application
    with app.app_context():
        hr = User(name="Recruiter Rita", email="rita@company.com", role="hr")
        hr.set_password("passrita123")
        db.session.add(hr)
        db.session.commit()

    # Log in as HR via HR login portal
    login_res = client.post("/hr/login", data={
        "email": "rita@company.com",
        "password": "passrita123"
    }, follow_redirects=True)
    assert login_res.status_code == 200
    assert b"Candidate Skill Verification Directory" in login_res.data

    # Candidate attempting HR portal
    client.get("/logout")
    client.post("/register", data={"name": "Normal Candidate", "email": "norm@example.com", "password": "normpass123"})
    client.post("/login", data={"email": "norm@example.com", "password": "normpass123"}, follow_redirects=True)

    block_res = client.get("/hr/dashboard", follow_redirects=True)
    assert block_res.status_code == 200
    assert b"Access denied" in block_res.data


def test_resume_upload_and_skill_extraction_pipeline(client):
    """Verify resume file upload, text extraction, and skill review."""
    # Register & Login
    client.post("/register", data={"name": "Resume User", "email": "resume@example.com", "password": "password123"})
    client.post("/login", data={"email": "resume@example.com", "password": "password123"}, follow_redirects=True)

    # Build DOCX resume
    doc = docx.Document()
    doc.add_heading("Resume User", 0)
    doc.add_paragraph("Email: resume@example.com | Phone: 9876543210")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, Java, MySQL, Flask, Docker, Git, REST API")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.E. Computer Science, 2024")
    
    file_bytes = io.BytesIO()
    doc.save(file_bytes)
    file_bytes.seek(0)

    # Upload
    upload_res = client.post(
        "/resume/upload",
        data={"resume_file": (file_bytes, "test_resume.docx")},
        content_type="multipart/form-data",
        follow_redirects=True
    )
    assert upload_res.status_code == 200
    assert b"AI Resume Analysis Review" in upload_res.data
    assert b"Python" in upload_res.data
    assert b"Java" in upload_res.data

    # Confirm claimed skills
    review_res = client.post("/resume/review", data={
        "skills_list": "Python, Java, MySQL, Flask, Docker, Git, REST API"
    }, follow_redirects=True)
    assert review_res.status_code == 200
    assert b"Claimed skills confirmed" in review_res.data
    assert b"Select Programming Language" in review_res.data


def test_phase1_mcq_assessment_and_answer_review(client):
    """Verify 7-question MCQ flow, proportional scoring, and answer review page."""
    # Register and login
    client.post("/register", data={"name": "Kavita Sen", "email": "kavita@example.com", "password": "password123"})
    client.post("/login", data={"email": "kavita@example.com", "password": "password123"}, follow_redirects=True)

    # Instructions page
    instr_res = client.get("/assessment/select-language")
    assert instr_res.status_code == 200
    assert b"Select Programming Language" in instr_res.data
    assert b"Python" in instr_res.data
    assert b"Java" in instr_res.data
    assert b"C Programming" in instr_res.data
    assert b"C++" in instr_res.data

    # Start Python test
    start_res = client.get("/assessment/start/python", follow_redirects=True)
    assert start_res.status_code == 200
    assert b"Python Skill Assessment" in start_res.data
    assert b"Question 1 of 7" in start_res.data

    # Extract attempt ID from response URL or database
    app = client.application
    with app.app_context():
        user = User.query.filter_by(email="kavita@example.com").first()
        attempt = Attempt.query.filter_by(user_id=user.id, status="in_progress").first()
        assert attempt is not None
        questions = attempt.get_selected_questions()
        assert len(questions) == 7

        # Form submission: 6 correct, 1 incorrect
        sub_data = {}
        for idx, q in enumerate(questions):
            if idx < 6:
                sub_data[f"q_{q.id}"] = q.correct_answer
            else:
                sub_data[f"q_{q.id}"] = "Incorrect Option Choice"
        att_id = attempt.id

    submit_res = client.post(f"/assessment/take/{att_id}", data=sub_data, follow_redirects=True)
    assert submit_res.status_code == 200
    assert b"Assessment submitted successfully" in submit_res.data
    assert b"85.71" in submit_res.data
    assert b"6" in submit_res.data  # 6 correct
    assert b"1" in submit_res.data  # 1 incorrect
    assert b"Complete Question-by-Question Answer Review" in submit_res.data
    assert b"Technical Explanation:" in submit_res.data
    assert b"Download PDF Assessment Report" in submit_res.data
    assert b"Download Verified Certificate" in submit_res.data


def test_candidate_privacy_and_hr_inspection(client):
    """Verify Candidate A cannot access Candidate B's report, while HR can inspect it."""
    app = client.application
    with app.app_context():
        cand_a = User(name="Alice", email="alice@priv.com", role="candidate")
        cand_a.set_password("pass123")
        cand_b = User(name="Bob", email="bob@priv.com", role="candidate")
        cand_b.set_password("pass123")
        hr_user = User(name="Recruiter", email="hr@priv.com", role="hr")
        hr_user.set_password("pass123")
        db.session.add_all([cand_a, cand_b, hr_user])
        db.session.commit()

        # Alice completes test
        att_a = start_or_resume_attempt(cand_a.id, "python")
        q_list = att_a.get_selected_questions()
        evaluate_assessment(att_a.id, {f"q_{q.id}": q.correct_answer for q in q_list})
        att_a_id = att_a.id
        rep_a_id = att_a.report.id

    # Bob attempts to view Alice's results -> 403
    client.post("/login", data={"email": "bob@priv.com", "password": "pass123"}, follow_redirects=True)
    res_b = client.get(f"/assessment/result/{att_a_id}")
    assert res_b.status_code == 403
    res_b_rep = client.get(f"/report/{rep_a_id}")
    assert res_b_rep.status_code == 403

    # HR inspects Alice's report -> 200 OK
    client.get("/logout")
    client.post("/login", data={"email": "hr@priv.com", "password": "pass123"}, follow_redirects=True)
    res_hr = client.get(f"/assessment/result/{att_a_id}")
    assert res_hr.status_code == 200
    res_hr_rep = client.get(f"/report/{rep_a_id}")
    assert res_hr_rep.status_code == 200
